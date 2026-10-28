"""Build V4 original diagrams, Word manuals and blank records.
Run using the Codex bundled Python; never changes previous versions.
"""
from pathlib import Path
import math, re, json, csv, html
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / 'assets'
ASSETS.mkdir(exist_ok=True)
INK, BLUE, RED, EARTH = '#20313b', '#24647d', '#b13c34', '#dfc8a7'

class Drawing:
    def __init__(self, title, h=850):
        self.w, self.h = 1600, h
        self.im = Image.new('RGB', (self.w,h), 'white')
        self.d = ImageDraw.Draw(self.im)
        self.svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="1600" height="{h}" viewBox="0 0 1600 {h}">', '<rect width="100%" height="100%" fill="white"/>']
        self.text((45,25), title, 36, True)
    def text(self, p, t, size=30, bold=False, fill=INK):
        ff='C:/Windows/Fonts/msyhbd.ttc' if bold else 'C:/Windows/Fonts/msyh.ttc'
        font=ImageFont.truetype(ff,size)
        for j,line in enumerate(t.split('\n')):
            x,y=p[0],p[1]+j*(size+12)
            self.d.text((x,y),line,font=font,fill=fill)
            self.svg.append(f'<text x="{x}" y="{y+size}" font-family="Microsoft YaHei,sans-serif" font-size="{size}" font-weight="{700 if bold else 400}" fill="{fill}">{html.escape(line)}</text>')
    def line(self, points, fill=INK, width=3):
        self.d.line(points, fill=fill, width=width)
        pts=' '.join(f'{x},{y}' for x,y in points)
        self.svg.append(f'<polyline points="{pts}" fill="none" stroke="{fill}" stroke-width="{width}"/>')
    def rect(self, box, fill='white', outline=INK, width=2):
        x1,y1,x2,y2=box
        self.d.rectangle(box,fill=fill,outline=outline,width=width)
        self.svg.append(f'<rect x="{x1}" y="{y1}" width="{x2-x1}" height="{y2-y1}" fill="{fill}" stroke="{outline}" stroke-width="{width}"/>')
    def circle(self, x,y,r=9, fill=RED):
        self.d.ellipse((x-r,y-r,x+r,y+r),fill=fill)
        self.svg.append(f'<circle cx="{x}" cy="{y}" r="{r}" fill="{fill}"/>')
    def arrow(self,a,b,fill=BLUE,both=False,width=4):
        self.line([a,b],fill,width)
        def head(q,p):
            th=math.atan2(p[1]-q[1],p[0]-q[0])
            pts=[p,(p[0]-15*math.cos(th-.45),p[1]-15*math.sin(th-.45)),(p[0]-15*math.cos(th+.45),p[1]-15*math.sin(th+.45)),p]
            self.line(pts,fill,width)
        head(a,b)
        if both:head(b,a)
    def save(self,name):
        self.im.save(ASSETS/(name+'.png'))
        (ASSETS/(name+'.svg')).write_text('\n'.join(self.svg+['</svg>']),encoding='utf-8')

def figures():
    f=Drawing('横向布纤的方向响应   理想几何算例不等于土中检测限',850)
    cases=[(60,'沿轴线张开','L = 100 mm  Δy = 10 μm','理想平均应变 100 με'),(570,'垂直轴线差动','L = 100 mm  Δs = 10 μm','理想平均应变 0.005 με'),(1080,'整个区域同动','两端同向同量移动','相对伸长为零') ]
    for x,title,sub,result in cases:
        f.text((x,125),title,32,True)
        f.text((x,180),sub,27)
        f.line([(x+25,320),(x+410,320)],'#a9b2b8',4)
        if x==60:
            f.line([(x+25,320),(x+440,320)],RED,6);f.arrow((x+360,260),(x+450,260));f.text((x+100,380),'沿横坡 y',28)
        elif x==570:
            f.line([(x+25,320),(x+410,375)],RED,6);f.arrow((x+410,290),(x+410,375));f.text((x+60,405),'沿 s 或 n   图中位移夸大',26)
        else:
            f.line([(x+25,375),(x+410,375)],RED,6);f.arrow((x+25,290),(x+25,375));f.arrow((x+410,290),(x+410,375));f.text((x+65,405),'两端之间无差动',28)
        f.text((x,485),result,29,True,RED)
    f.text((60,590),'局部弯曲转换件：将两锚点的顺坡或法向差动转换为表面轴向应变。',32)
    f.text((60,650),'仍须验证土锚滑移、转动、温漂与构件加筋效应；不能消除整块同动盲区。',30)
    f.text((60,730),'灰线为初始位置；红线为变形示意。主光路保持横向，示意位移非实际比例。',27)
    f.save('F01_方向与盲区')

    f=Drawing('B土箱横向阵列平面图   净尺寸 s600 × y800 mm',980)
    x0,y0,k=170,170,1.0
    f.rect((x0,y0,x0+800,y0+600),'#faf4eb',BLUE,4)
    f.arrow((x0,125),(x0+800,125),both=True);f.text((480,80),'横坡 y  800',30)
    f.arrow((115,y0),(115,y0+600),both=True);f.text((30,440),'s\n600',27)
    for i,s in enumerate((150,300,450),1):
        yy=y0+s
        f.line([(x0-65,yy),(x0+780,yy)],BLUE,4)
        f.text((990,yy-17),f'H{i}  s={s}',27)
        for j,y in enumerate((200,400,600),1):
            x=x0+y
            f.rect((x-90,yy-25,x+90,yy+25),'#ead4b5',RED,2)
            f.circle(x,yy);f.text((x-49,yy-66),f'H{i}-{j}',25)
        f.circle(x0+100,yy-30,8,BLUE);f.text((x0+48,yy-72),f'T{i}',24,False,BLUE)
    for y in (200,400,600):f.text((x0+y-30,805),str(y),28)
    f.text((1230,170),'节点最大占位\n180 × 50\n埋深 d = 80',27)
    f.text((1230,340),'每线三个节点\n中心间距 200\n侧边余量 110',27)
    f.text((1230,515),'左侧分路引出\n箱外保护汇线\n不纵穿土体',27)
    f.text((80,875),'A：9应变 + 3温度 = 12栅     B单方向：18应变 + 3温度 = 21栅',30)
    f.text((80,925),'温度点独立松套；节点之间保留柔性余量。各配置分别重建对比。',27)
    f.save('F02_横向阵列')

    f=Drawing('B转换节点尺寸与传力   横向两锚点差动弯曲',1010)
    x0,y,k=200,275,6
    f.rect((x0-60,y-65,x0+180,y+65),'#d4e0e6',BLUE,3)
    f.rect((x0+780,y-65,x0+1020,y+65),'#d4e0e6',BLUE,3)
    f.rect((x0,y-16,x0+960,y+16),'#b8c6cd',INK,3)
    f.rect((x0,y-25,x0+120,y+25),'#a2b5bf',INK,3)
    f.rect((x0+840,y-25,x0+960,y+25),'#a2b5bf',INK,3)
    gx=x0+120+15*k
    f.rect((gx-15,y-26,gx+15,y-18),RED,RED)
    f.rect((gx-15,y+18,gx+15,y+26),RED,RED)
    f.arrow((x0,135),(x0+960,135),both=True);f.text((590,86),'薄片总长 160',29)
    f.arrow((x0+120,390),(x0+840,390),both=True);f.text((565,410),'有效跨长 120',29)
    f.text((x0-40,465),'夹持 20',28);f.text((x0+810,465),'夹持 20',28)
    f.text((1260,225),'局部锚片\n约 40 × 40\n土中不接箱壁',27)
    f.text((80,550),'FBG两面同位：距有效跨根部 a = 15；栅长候选5。',30)
    f.text((80,605),'SUS301弹性片：宽10；厚0.3 / 0.5 / 0.8三组；柔性护罩不跨接传力。',29)
    f.line([(gx,y+30),(gx,525),(380,525)],RED,2)
    f.text((80,680),'测法向 n 差动：宽面平行土面。测顺坡 s 差动：薄片绕 y 轴转90°。',29)
    f.text((80,740),'t = 0.5时理想响应 0.0781 με/μm；10 μm仅约0.781 με，须实测噪声。',29,True)
    f.text((80,805),'感测段 → 渐变柔性过渡 → 松套连接 → 箱外铠装或可拆保护管',30,False,BLUE)
    f.text((80,870),'固定导向梁仅为计算假设。锚片转动和拔滑必须独立量测。',29)
    f.text((80,930),'尺寸单位 mm；仅作样件制造讨论，夹具、胶层及限位在样件标定后冻结。',26)
    f.save('F03_转换节点')

    f=Drawing('土箱横剖面 y向   活动底板输入不等于土体内部位移',1010)
    x0,y0,k=120,170,1.1
    bottom=y0+330;surface=bottom-176
    f.rect((x0,y0,x0+880,bottom),'white',BLUE,5)
    for j in range(8):
        yy=bottom-(j+1)*22
        f.rect((x0+4,yy,x0+876,yy+22),'#dec7a6' if j%2 else '#ead6b8','#d3b994',1)
    f.line([(x0+100,bottom-88),(x0+780,bottom-88)],RED,5)
    for xx in (x0+220,x0+440,x0+660):f.circle(xx,bottom-88,8)
    f.rect((x0,bottom+5,x0+880,bottom+15),'#b8c9d2',BLUE,2)
    f.rect((x0,bottom+85,x0+880,bottom+95),'#b8c9d2',BLUE,2)
    f.rect((x0+30,bottom+95,x0+850,bottom+115),'#7c95a2',BLUE,2)
    f.rect((x0+462,bottom-3,x0+572,bottom+15),'#c77f60',RED,3)
    f.arrow((x0+517,bottom+20),(x0+517,bottom+70),RED)
    f.text((x0+20,bottom+30),'干燥机械层  高度按驱动器确定',24)
    f.text((1050,185),'净高 300\n土厚 160\n8层 × 20',29)
    f.text((1050,340),'d从土面向下\n红线默认 d80\n另次比较 d40/120',28)
    f.text((80,620),'304底板3 + 连续网格承托；外架方管30 × 30 × 3候选。',29)
    f.text((80,675),'PC观察壁8 + 外压框；接缝EPDM密封；执行器与相机另架支承。',29)
    f.text((80,745),'M1 双盘差动\n沿y张开 / 沿s剪切',28,True)
    f.text((590,745),'M2 局部下沉板\n100 × 100 支承卸载',28,True)
    f.text((1100,745),'M3 局部托盘\n200 × 200 缓慢错动',28,True)
    f.text((80,865),'示意图未按板厚放大比例绘制。每次只安装一种模块，底板连续性随接口复核。',26)
    f.text((80,920),'水平土箱用于机制与耦合验证；坡槽阶段另定坡角和几何，保留V3 CAD作历史参考。',26)
    f.save('F04_土箱与模块')

    f=Drawing('评价路径与独立事件时标   示意时序不代表已获提前量',850)
    labels=[('台架标定','独立位移与电测\n灵敏度 噪声 温漂'),('土中传力','土标记与锚片滑移\n干湿和有无纤对照'),('病害检出','同误报预算比较\n检出概率及提前量')]
    for j,(title,body) in enumerate(labels):
        x=80+j*510
        f.text((x,130),title,34,True);f.text((x,195),body,30)
        if j<2:f.arrow((x+350,208),(x+475,208))
    f.arrow((100,480),(1490,480))
    ev=[(130,'t0','实际加载'),(420,'tinternal','内部参考'),(670,'talarm','因果报警'),(930,'tvisible','表面可见'),(1270,'tfailure','宏观破坏')]
    for x,label,body in ev:
        f.circle(x,480,8,RED if label=='talarm' else BLUE)
        f.text((x-48,380),label,27,True);f.text((x-55,525),body,27)
    f.text((80,620),'分别计算：表面可见时刻 − 报警时刻；持续运动时刻 − 报警时刻；破坏时刻 − 报警时刻。',25)
    f.text((80,680),'各事件顺序可变，负提前量照实记录；无破坏时不填写破坏提前量。',29)
    f.text((80,745),'训练测试按独立重建分组；30 h零误报的泊松95%上界仍约0.1次/h。',28)
    f.save('F05_评价与时标')

    f=Drawing('病害模块与锚点位置   同一土箱分次更换模块',1010)
    for col,(name,center,width) in enumerate([('M1 双盘差动',400,300),('M2 局部下沉',470,100),('M3 局部错动',500,200)]):
        ox=70+col*510;oy=180;scale=.85
        f.text((ox,120),name,32,True)
        f.rect((ox,oy,ox+425,oy+280),'#faf4eb',BLUE,3)
        yy=oy+140
        f.line([(ox,yy),(ox+425,yy)],BLUE,3)
        def px(y):return ox+(y-200)*scale
        if col==0:
            f.rect((px(250),yy-85,px(400),yy+85),'#d5e4eb',BLUE,2)
            f.rect((px(400),yy-85,px(550),yy+85),'#dfb99b',RED,2)
            f.arrow((px(480),yy+110),(px(530),yy+110),RED)
        else:
            half=width/2
            f.rect((px(center-half),yy-half*scale,px(center+half),yy+half*scale),'#dfb99b',RED,3)
            if col==2:f.arrow((px(520),yy+110),(px(520),yy+160),RED)
        for ya in (330,470):
            f.rect((px(ya-20),yy-17,px(ya+20),yy+17),'#8aa3ad',INK,2)
            f.text((px(ya)-30,oy+295),str(ya),25)
        f.line([(px(330),yy),(px(470),yy)],RED,5)
        f.text((ox,oy+350),'H2-2锚点 y330 / y470',26)
        if col==0:f.text((ox,oy+410),'分界 y400\n两锚点位于不同承载盘',27)
        elif col==1:f.text((ox,oy+410),'模块中心 y470\n右锚上方卸载 左锚参考',27)
        else:f.text((ox,oy+410),'模块边缘 y400 / y600\n同时考察H2-2和H2-3',27)
    f.text((70,780),'图示为s300测线局部平面；y显示200—700 mm。灰色为局部锚片，红线为转换梁。',27)
    f.text((70,840),'底部差动可能经土拱和剪切扩散而改变；d80处两锚点的相对运动必须独立测得。',28)
    f.text((70,905),'M2避免对称下沉只使梁跨中受土压的歧义；若变形区偏移，按土内参考调整测线。',27)
    f.save('F06_模块与锚点')

    f=Drawing('单端夹具与锚片装配   两端分别连接 不跨接有效弯曲段',1000)
    f.text((70,110),'单端俯视  尺寸单位 mm',30,True)
    f.rect((100,205,500,605),'#dce8ec',BLUE,3)
    f.rect((200,285,400,525),'#9db5c1',INK,3)
    f.rect((200,355,790,455),'#bbcbd2',INK,3)
    f.rect((200,285,400,525),'#9db5c1',INK,3)
    for cy in (325,485):f.circle(300,cy,17,'white');f.circle(300,cy,5,INK)
    f.arrow((100,165),(500,165),both=True);f.text((220,175),'锚片长40',26)
    f.arrow((200,650),(400,650),both=True);f.text((210,680),'夹长20',26)
    f.text((540,255),'两颗M3\n在片宽外侧\n距中线 ±8',27)
    f.text((535,480),'弹性片宽10\n不得穿孔夹栅\n夹块宽24',27)
    f.text((900,115),'夹持剖面',30,True)
    f.rect((940,220,1400,280),'#9db5c1',INK,3)
    f.rect((940,295,1460,311),'#bbcbd2',INK,2)
    f.rect((940,325,1400,385),'#9db5c1',INK,3)
    f.text((950,420),'上夹块厚4  下夹块厚4\n中间弹性片厚度按试验组\n示意留缝 实际夹紧接触',27)
    f.text((900,605),'本端角座连接本端锚片\n法向试验锚片平放\n顺坡试验锚片转入y-n面',28)
    f.text((70,790),'先台架测夹持滑移与转角，再测土锚拔滑；紧固扭矩与孔位配合由样件确定。',29)
    f.text((70,855),'锚片40 × 40 × 1、夹块20 × 24 × 4为制造讨论值；图示非签审加工图。',29)
    f.text((70,920),'光纤尾纤沿长向引出，避免在10 mm片宽内回折；连接弯曲半径按厂家数据。',27)
    f.save('F07_夹具与锚片')

def text_runs(p, text):
    for part in re.split(r'(https://[^\s。]+)', text.replace('**','')):
        if part.startswith('https://'):
            link=OxmlElement('w:hyperlink')
            rid=p.part.relate_to(part,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
            link.set(qn('r:id'),rid)
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'245A75');pr.append(col);r.append(pr)
            t=OxmlElement('w:t');t.text=part;r.append(t);link.append(r);p._p.append(link)
        else:p.add_run(part)

def setup(doc):
    sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
    sec.left_margin=Inches(.8);sec.right_margin=Inches(.8);sec.footer_distance=Inches(.28)
    for name,sz in [('Normal',11),('Title',24),('Subtitle',11),('Heading 1',17),('Heading 2',13),('Heading 3',11),('Caption',10)]:
        st=doc.styles[name];st.font.name='Microsoft YaHei';st.font.size=Pt(sz);st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        st.paragraph_format.space_after=Pt(6 if name=='Normal' else 9)
        st.paragraph_format.line_spacing=1.15
        if name.startswith('Heading'):st.paragraph_format.keep_with_next=True
        ppr=st.element.find(qn('w:pPr'))
        if ppr is not None:
            for border in list(ppr.findall(qn('w:pBdr'))):ppr.remove(border)
    doc.styles['Normal'].paragraph_format.widow_control=True
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    foot.add_run('V4.0  |  2026-10-28  |  ').font.size=Pt(9)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    doc.core_properties.author='桥梁边坡光纤监测项目组'
    doc.core_properties.last_modified_by='桥梁边坡光纤监测项目组'

def table(doc, rows):
    rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',c) for c in r)]
    n=len(rows[0]); weights={2:[.3,.7],3:[.23,.43,.34],4:[.2,.27,.27,.26]}.get(n,[1/n]*n)
    t=doc.add_table(rows=0,cols=n);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,w in zip(t.columns,weights):c.width=Inches(6.9*w)
    b=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
    t._tbl.tblPr.append(b)
    for i,row in enumerate(rows):
        cells=t.add_row().cells;trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,txt) in enumerate(zip(cells,row)):
            cell.width=Inches(6.9*weights[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');m.append(e)
            tcpr.append(m);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE8EE' if i==0 else 'FFFFFF');tcpr.append(sh)
            p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.12
            if i==0:p.paragraph_format.keep_with_next=True
            if len(txt)<14:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(txt);r.font.size=Pt(10);r.bold=(i==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def build(path):
    doc=Document();setup(doc);source=path.read_text(encoding='utf-8')
    if path.name.startswith('02_'):
        block='![横向布纤方向性](assets/F01_方向与盲区.png)\n\n图1 横向布纤对不同位移方向的响应路径及盲区\n\n'
        source=source.replace(block,'')
        source=source.replace('## 1 坐标与方向可观测性',block+'## 1 坐标与方向可观测性')
        intro='节点使用横向弹性薄片连接两个局部锚片；在薄片两面贴FBG，利用差动弯曲产生一拉一压。薄片长轴沿y，薄片的薄厚方向朝目标位移方向。测法向沉陷时薄片宽面平行坡面；测顺坡剪切时薄片转90°，宽面沿y-n平面。更改的是节点姿态，主光缆仍横向。'
        source=source.replace(intro+'\n\n','')
        caption='图3 双面FBG弯曲件的候选尺寸与引线分区'
        source=source.replace(caption,caption+'\n\n'+intro)
    lines=source.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append([x.strip() for x in lines[i].strip().strip('|').split('|')]);i+=1
            table(doc,rows);continue
        if line.startswith('# '):doc.add_paragraph(line[2:],style='Title')
        elif line.startswith('### '):doc.add_paragraph(line[4:],style='Heading 2')
        elif line.startswith('## '):doc.add_paragraph(line[3:],style='Heading 1')
        elif line.startswith('!['):
            m=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',line);p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            imagepath=ROOT/m[2]
            with Image.open(imagepath) as im:iw,ih=im.size
            shape=p.add_run().add_picture(str(imagepath),width=Inches(min(6.9,4.4*iw/ih)));shape._inline.docPr.set('descr',m[1])
        elif re.match(r'^图\d+',line):
            p=doc.add_paragraph(line,style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        else:
            p=doc.add_paragraph();text_runs(p,line)
        i+=1
    doc.save(path.with_suffix('.docx'))
    return {'file':path.with_suffix('.docx').name,'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'figures':len(doc.inline_shapes)}

def records():
    p=ROOT/'templates';p.mkdir(exist_ok=True)
    specs={
      '01_试验登记.csv':['trial_id','rebuild_id','material_batch','specimen_id','configuration','mechanism','date_time_with_timezone','operator','calibration_id','instrument_id','firmware','git_commit','planned_input','measured_input_file','raw_wavelength_file','reference_file','video_file','split_group','excluded_reason'],
      '02_方向标定.csv':['calibration_id','specimen_id','cycle_id','axis','angle_deg','command_um','measured_um','reference_expanded_uncertainty_um','surface_A_ue','surface_B_ue','half_difference_ue','temperature_C','force_N','bandwidth_Hz','hold_s','zero_drift_ue','loading_or_unloading','quality_flag'],
      '03_事件与提前量.csv':['trial_id','event','time_lower_s','time_upper_s','independent_source','criterion_version','synchronization_uncertainty_s','adjudicator','alarm_algorithm_version','endpoint','lead_lower_s','lead_upper_s','observation_end_s','censored_reason'],
      '04_检测与误报评价.csv':['analysis_id','test_group','algorithm_version','threshold_version','bandwidth_Hz','normal_effective_h','false_alarm_events','false_alarm_rate_per_h','false_alarm_upper95_per_h','independent_trials','detected_trials','Pd','Pd_lower95','Pd_upper95','minimum_input_um','localization_error_mm','missing_channels','confidence_method'],
      '05_铺土与耦合.csv':['trial_id','layer_from_bottom','thickness_mm','net_volume_m3','dry_sand_kg','dry_kaolin_kg','sand_initial_water_ratio','kaolin_initial_water_ratio','added_water_kg','wet_mass_kg','measured_water_ratio','measured_dry_density_kg_m3','compaction_method','node_id','s_mm','y_mm','d_mm','anchor_slip_um','prebury_spectrum_file','postbury_spectrum_file','remarks']
    }
    for name,fields in specs.items():
        if not (p/name).exists():
            with (p/name).open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerow(fields)
    cfg={'status':'planned_not_measured','version':'V4.0','coordinates':{'s':'downslope/loading direction','y':'transverse','n':'outward normal','d':'depth below local soil surface'},'box_mm':{'s':600,'y':800,'height':300,'soil_thickness':160},'line_s_mm':[150,300,450],'node_y_mm':[200,400,600],'depth_mm':80,'configuration_counts':{'A':12,'B_single_direction':21,'B_two_direction':39},'soil_candidates_dry_mass':{'A0':{'sand':1.0,'kaolin':0.0},'A5':{'sand':.95,'kaolin':.05},'A10':{'sand':.9,'kaolin':.1}},'candidate_water_ratio':[.02,.04,.06],'actual_dry_density_kg_m3':None,'actual_interrogator':None,'measured_noise_ue':None,'validated_detection_limit_um':None,'alarm_rule':None,'test_split_rule':'group by independent rebuild and batch before windowing'}
    if not (p/'planned_configuration.json').exists():
        (p/'planned_configuration.json').write_text(json.dumps(cfg,ensure_ascii=False,indent=2),encoding='utf-8')
    (p/'README.md').write_text('# V4试验记录空表\n\n这些文件仅定义字段，没有预填实测成绩。位移、应变、时间和质量的单位写在列名中。时间用带时区记录；事件时间以同一试验单调时钟为基准。所有重复循环仍属于同一重建组。\n\n配方含水率采用干基小数，例如0.04代表4%。配置JSON中的null表示未取得数据，不表示零。阶段、参数和标签请在试验前冻结并记录版本。\n',encoding='utf-8')

def calculations():
    L=100.0;delta=.01
    # Rationalized exact expression avoids subtraction cancellation.
    perp=(delta/L)**2/(math.sqrt(1+(delta/L)**2)+1)*1e6
    beam=3*.5/(120**2)*(1-2*15/120)*1000
    vol=.6*.8*.16;dry=vol*1500;porosity=1-1500/2650
    result={'status':'analytical_design_examples_not_measurements','transverse_perpendicular_10um_ue':perp,'axial_10um_ue':delta/L*1e6,'ideal_beam_ue_per_um':beam,'ideal_beam_10um_ue':beam*10,'box_soil_m3':vol,'dry_mass_example_kg':dry,'wet_mass_4percent_kg':dry*1.04,'one_layer_dry_kg':.6*.8*.02*1500,'one_layer_wet_4percent_kg':.6*.8*.02*1500*1.04,'saturated_pore_water_example_kg':vol*porosity*1000,'saturated_soil_water_example_kg':dry+vol*porosity*1000,'20mmh_20min_litres':.48*20/3,'zero_alarm_30h_poisson_upper95_per_h':-math.log(.05)/30,'fbg_counts':[9+3,9*2+3,9*4+3],'node_side_clearance_mm':200-90}
    assert math.isclose(perp,.005,rel_tol=1e-6)
    assert math.isclose(beam,.078125) and math.isclose(dry,115.2)
    assert result['fbg_counts']==[12,21,39]
    (ROOT/'计算复核.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return result

if __name__=='__main__':
    figures();records();checks=calculations()
    docs=[build(p) for p in sorted(ROOT.glob('0[1-4]_*.md'))]
    print(json.dumps({'documents':docs,'calculations':checks},ensure_ascii=False,indent=2))
