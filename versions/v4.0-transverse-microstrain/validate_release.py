"""Validate V4 artifacts and design arithmetic, not experimental performance."""
from pathlib import Path
from urllib.parse import unquote
import csv
import hashlib
import json
import math
import re
import zipfile
from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[1]

def require(condition, message):
    if not condition:
        raise AssertionError(message)

body = (ROOT/'01_病害与文献证据.md').read_text(encoding='utf-8')
source_ids = re.findall(r'^### (R[0-9]{2}) ', body, re.M)
defect_ids = re.findall('^' + re.escape('| ') + r'(D[0-9]{2}) ', body, re.M)
require(source_ids == [f'R{i:02}' for i in range(1,20)], '19 source headings')
require(defect_ids == [f'D{i:02}' for i in range(1,29)], '28 defect/process rows')
markdown = list(ROOT.glob('*.md')) + list((ROOT/'templates').glob('*.md')) + [REPO/'README.md']
link_count = 0
link_pattern = r'!?' + re.escape('[') + r'.*?' + re.escape('](') + r'([^)]+)' + re.escape(')')
for path in markdown:
    for match in re.finditer(link_pattern, path.read_text(encoding='utf-8')):
        target = match.group(1)
        if target.startswith(('https:', 'http:', '#', 'mailto:')):
            continue
        local = unquote(target.split('#', 1)[0].strip('<>'))
        require((path.parent/local).exists(), f'Missing link {path.name}: {target}')
        link_count += 1
require(link_count > 30, 'Expected authored local links to be checked')

docs = sorted(ROOT.glob('0[1-4]_*.docx'))
require(len(docs) == 4, 'Four Word documents')
doc_report = []
for path in docs:
    with zipfile.ZipFile(path) as archive:
        require(archive.testzip() is None, f'Corrupt DOCX: {path.name}')
    doc = Document(path)
    content = chr(10).join(p.text for p in doc.paragraphs)
    require('2026年10月28日' in content and '研究设计版' in content, f'Wrong edition: {path.name}')
    require('最终复核待完成' not in content, f'Stale header: {path.name}')
    doc_report.append({'file':path.name, 'paragraphs':len(doc.paragraphs), 'tables':len(doc.tables), 'figures':len(doc.inline_shapes)})
pngs = sorted((ROOT/'assets').glob('F*.png'))
svgs = sorted((ROOT/'assets').glob('F*.svg'))
require(len(pngs) == len(svgs) == 7, 'Seven PNG and SVG diagrams')
require({p.stem for p in pngs} == {p.stem for p in svgs}, 'Matching diagram stems')
for path in pngs:
    with Image.open(path) as im:
        im.verify()
for path in svgs:
    require('<svg' in path.read_text(encoding='utf-8'), f'SVG root missing: {path.name}')

csvs = sorted((ROOT/'templates').glob('*.csv'))
require(len(csvs) == 5, 'Five CSV templates')
for path in csvs:
    with path.open(encoding='utf-8-sig', newline='') as stream:
        rows = list(csv.reader(stream))
    require(len(rows) == 1 and all(rows[0]), f'Template is not header-only: {path.name}')
    require(len(rows[0]) == len(set(rows[0])), f'Duplicate CSV fields: {path.name}')
cfg = json.loads((ROOT/'templates/planned_configuration.json').read_text(encoding='utf-8'))
require(cfg['status'] == 'planned_not_measured', 'Configuration status')
require(cfg['box_mm'] == {'s':600,'y':800,'height':300,'soil_thickness':160}, 'Box geometry')
require(cfg['line_s_mm'] == [150,300,450] and cfg['node_y_mm'] == [200,400,600], 'Node coordinates')
require(cfg['depth_mm'] == 80, 'Planned embedment')
require(list(cfg['configuration_counts'].values()) == [12,21,39], 'FBG counts')
for name in ('actual_dry_density_kg_m3','actual_interrogator','measured_noise_ue','validated_detection_limit_um','alarm_rule'):
    require(cfg[name] is None, f'Unmeasured field contains result: {name}')
for blend in cfg['soil_candidates_dry_mass'].values():
    require(math.isclose(sum(blend.values()),1.0), 'Dry blend fractions')
calc = json.loads((ROOT/'计算复核.json').read_text(encoding='utf-8'))
expected = {'transverse_perpendicular_10um_ue':.005,'axial_10um_ue':100,'ideal_beam_ue_per_um':.078125,'ideal_beam_10um_ue':.78125,'box_soil_m3':.0768,'dry_mass_example_kg':115.2,'wet_mass_4percent_kg':119.808,'one_layer_dry_kg':14.4,'one_layer_wet_4percent_kg':14.976,'20mmh_20min_litres':3.2,'zero_alarm_30h_poisson_upper95_per_h':-math.log(.05)/30}
for key,value in expected.items():
    require(math.isclose(calc[key],value,rel_tol=1e-7), f'Calculation mismatch: {key}')
require(math.isclose(9.5*1.01+.5*1.02+.295,10*1.04), 'Initial moisture correction')
require(math.isclose(1500*1.04*9.81*.08/1000,1.224288), 'Self-weight total stress')
require(330 < 420 and 420 < 470 < 520, 'M2 anchor/moving-window geometry')
require(330 < 400 < 470 and 530 < 600 < 670, 'M3 boundaries cross two nodes')
manifest_path = ROOT/'_qa/render_manifest.json'
rendered = None
if manifest_path.exists():
    rendered = json.loads(manifest_path.read_text(encoding='utf-8'))
    require([r['pages'] for r in rendered] == [9,8,9,8], 'Final rendering page count')
    for item in rendered:
        require(all(Path(p).exists() for p in item['page_files']), 'Missing rendered page')
        require(min(item['text_chars_per_page']) > 100, 'Unexpected blank or image-only page')
report = {'status':'PASS','source_entries':len(source_ids),'defect_process_entries':len(defect_ids),'documents':doc_report,'diagrams':len(pngs),'blank_csv_templates':len(csvs),'rendered_pages':sum(r['pages'] for r in rendered) if rendered else None,'docx_sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in docs},'scope':'Document integrity and analytical design examples only; empirical performance unverified'}
qa = ROOT/'_qa'
qa.mkdir(exist_ok=True)
(qa/'release_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
